#!/usr/bin/env python3
"""
enricher.py - Enrich an Excel spreadsheet of business contacts using Apify actors.

Pipeline per incomplete contact:
    1. LinkedIn people search  (get-leads/linkedin-scraper)       -> email, LinkedIn URL
    2. Company intelligence    (foxlabs/owler-intelligence,
                                fallback automation-lab/owler-...)  -> website, address,
                                                                      revenue, employees
    3. LinkedIn phone scraper  (api-empire/linkedin-profile-phone-number-scraper)
                                                                    -> business phone

Features: multi-key round-robin rotation with budget checks, exhausted-key
tracking, resumable progress (progress.json), periodic checkpoints, and a
strict "only fill empty cells" merge policy.

Usage:
    python enricher.py input.xlsx [--dry-run] [--test N] [--resume] [--status]
                                  [--skip-phones] [--keys keys.json] [--delay S]

Dependencies: requests, openpyxl (everything else is standard library).
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from datetime import datetime, timedelta

import openpyxl
import requests
from openpyxl.styles import PatternFill

# ════════════════════════════════════════════════════════════════════════════
# Constants (immutable configuration - no mutable global state)
# ════════════════════════════════════════════════════════════════════════════

API_BASE = "https://api.apify.com/v2"

DEFAULT_ACTORS = {
    "linkedin_search": "get-leads/linkedin-scraper",
    "company_primary": "foxlabs/owler-intelligence",
    "company_fallback": "automation-lab/owler-company-intelligence-scraper",
    "phone": "api-empire/linkedin-profile-phone-number-scraper",
}

# Input for the phone actor. "{profile_url}" and "{li_at}" are substituted at
# runtime. Override with "phone_actor_input" in keys.json if the actor's input
# schema differs.
DEFAULT_PHONE_INPUT_TEMPLATE = {
    "profileUrls": ["{profile_url}"],
    "cookie": "{li_at}",
}

# Internal field name -> expected Excel header.
FIELD_HEADERS = {
    "name": "Name",
    "email": "Email",
    "company": "Company",
    "street": "Street Address",
    "city": "City",
    "state": "State",
    "zip": "Zip Code",
    "phone": "Business Phone",
    "website": "Website",
    "revenue": "Annual Revenue",
    "employees": "Number of Employees",
    "completed": "Completed?",
}
COLUMN_ORDER = list(FIELD_HEADERS.keys())

# A row is "incomplete" if any of these is empty.
TARGET_FIELDS = ["email", "street", "zip", "phone", "website", "revenue", "employees"]
COMPANY_FIELDS = ["website", "street", "city", "state", "zip", "revenue", "employees"]

FIELD_STAT = {
    "email": "emails_found",
    "phone": "phones_found",
    "website": "websites_found",
    "revenue": "revenue_found",
    "employees": "employees_found",
}

TERMINAL_RUN_STATUSES = {"SUCCEEDED", "FAILED", "ABORTED", "TIMED-OUT"}
EXHAUSTION_RETRY_AFTER = timedelta(hours=24)
AUTO_ENRICHED_LABEL = "Auto-enriched"
HIGHLIGHT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

COMPANY_STOPWORDS = {
    "the", "of", "and", "a", "an", "llc", "inc", "incorporated", "corp", "corporation",
    "co", "company", "ltd", "limited", "llp", "lp", "pllc", "pc", "pa", "plc", "group",
    "holdings", "services", "solutions", "partners", "associates", "international",
}
NAME_SUFFIXES = {"jr", "sr", "ii", "iii", "iv", "md", "phd", "cpa", "esq", "mba", "do", "rn", "dr", "mr", "ms", "mrs"}

US_STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA",
    "colorado": "CO", "connecticut": "CT", "delaware": "DE", "district of columbia": "DC",
    "florida": "FL", "georgia": "GA", "hawaii": "HI", "idaho": "ID", "illinois": "IL",
    "indiana": "IN", "iowa": "IA", "kansas": "KS", "kentucky": "KY", "louisiana": "LA",
    "maine": "ME", "maryland": "MD", "massachusetts": "MA", "michigan": "MI", "minnesota": "MN",
    "mississippi": "MS", "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
    "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY",
    "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK", "oregon": "OR",
    "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC", "south dakota": "SD",
    "tennessee": "TN", "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA",
    "washington": "WA", "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
}


# ════════════════════════════════════════════════════════════════════════════
# Small pure helpers
# ════════════════════════════════════════════════════════════════════════════

def now_iso() -> str:
    """Return the current local time as an ISO-8601 string (seconds precision)."""
    return datetime.now().isoformat(timespec="seconds")


def parse_iso(value: str | None) -> datetime | None:
    """Parse an ISO-8601 timestamp, returning None if it is missing or invalid."""
    if not value:
        return None
    try:
        return datetime.fromisoformat(str(value))
    except ValueError:
        return None


def is_empty(value: object) -> bool:
    """True when a cell value counts as empty (None or whitespace-only string)."""
    return value is None or (isinstance(value, str) and value.strip() == "")


def tokenize(text: object) -> list[str]:
    """Lower-case text and split it into alphanumeric tokens."""
    cleaned = "".join(ch.lower() if ch.isalnum() else " " for ch in str(text or ""))
    return cleaned.split()


def company_tokens(name: object) -> set[str]:
    """Significant tokens of a company name (legal suffixes and filler removed)."""
    return {t for t in tokenize(name) if t not in COMPANY_STOPWORDS}


def company_cache_key(name: object) -> str:
    """Normalised key used to cache company lookups."""
    return " ".join(tokenize(name))


def owler_slug(company: str) -> str:
    """Best-effort Owler URL slug: lower-case alphanumerics, legal suffixes dropped."""
    legal = {"llc", "inc", "incorporated", "corp", "corporation", "ltd", "limited", "llp", "pllc", "pc", "plc", "co"}
    tokens = [t for t in tokenize(company) if t not in legal]
    return "".join(tokens)


def format_revenue(amount: float) -> str:
    """Format a USD amount in the sheet's style, e.g. 4200000 -> '$4.2M'."""
    for divisor, suffix in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
        if amount >= divisor:
            text = f"{amount / divisor:.1f}".rstrip("0").rstrip(".")
            return f"${text}{suffix}"
    return f"${int(amount)}"


def normalize_phone(raw: object) -> str | None:
    """Normalise a phone number; US 10-digit numbers become '(804) 334-8558'."""
    if raw is None:
        return None
    text = str(raw).strip()
    digits = "".join(ch for ch in text if ch.isdigit())
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    if 7 <= len(digits) <= 15:
        return text
    return None


def is_valid_email(value: object) -> bool:
    """Light sanity check for an email address."""
    if not isinstance(value, str):
        return False
    value = value.strip()
    if " " in value or value.count("@") != 1:
        return False
    local, domain = value.split("@")
    return bool(local) and "." in domain and not domain.startswith(".") and not domain.endswith(".")


def to_number(value: object) -> float | None:
    """Convert ints/floats/strings like '8,200' to a float, else None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        digits = "".join(ch for ch in value if ch.isdigit() or ch == ".")
        try:
            return float(digits) if digits else None
        except ValueError:
            return None
    return None


def state_abbrev(state: object) -> str | None:
    """Convert a US state name to its 2-letter code; pass through existing codes."""
    if is_empty(state):
        return None
    text = str(state).strip()
    if len(text) == 2 and text.isalpha():
        return text.upper()
    return US_STATES.get(text.lower(), text)


def mask_secrets(data: object) -> object:
    """Return a deep copy of data with cookie/token values masked for logging."""
    if isinstance(data, dict):
        masked = {}
        for key, value in data.items():
            lowered = str(key).lower()
            if any(word in lowered for word in ("cookie", "li_at", "token", "password")):
                masked[key] = "***"
            else:
                masked[key] = mask_secrets(value)
        return masked
    if isinstance(data, list):
        return [mask_secrets(item) for item in data]
    return data


def truncate(text: str, limit: int = 1500) -> str:
    """Shorten long strings for the log file."""
    return text if len(text) <= limit else text[:limit] + f"... [{len(text) - limit} more chars]"


def atomic_write_json(path: str, data: dict) -> None:
    """Write JSON to path atomically (temp file + rename) so crashes never corrupt it."""
    tmp_path = path + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as handle:
        json.dump(data, handle, indent=2, ensure_ascii=False)
    os.replace(tmp_path, path)


def load_json_file(path: str) -> dict:
    """Load a JSON file, ignoring full-line '//' comments and '_comment*' keys."""
    with open(path, "r", encoding="utf-8") as handle:
        lines = [line for line in handle if not line.lstrip().startswith("//")]
    data = json.loads("".join(lines))

    def strip_comments(obj: object) -> object:
        if isinstance(obj, dict):
            return {k: strip_comments(v) for k, v in obj.items() if not str(k).startswith("_")}
        if isinstance(obj, list):
            return [strip_comments(v) for v in obj]
        return obj

    return strip_comments(data)


# ════════════════════════════════════════════════════════════════════════════
# Logging
# ════════════════════════════════════════════════════════════════════════════

class RunLogger:
    """Writes a detailed timestamped log file and short console messages."""

    def __init__(self, log_path: str | None) -> None:
        """Open the log file (if a path is given) in line-buffered append mode."""
        self.log_path = log_path
        self._handle = open(log_path, "a", encoding="utf-8", buffering=1) if log_path else None

    def _write(self, level: str, message: str) -> None:
        """Append one line to the log file."""
        if self._handle:
            self._handle.write(f"{now_iso()} [{level}] {message}\n")

    def debug(self, message: str) -> None:
        """Log-file only message (API calls, payloads, raw results)."""
        self._write("DEBUG", message)

    def info(self, message: str, console: bool = False) -> None:
        """Informational message; optionally echoed to the console."""
        self._write("INFO", message)
        if console:
            print(message, flush=True)

    def warn(self, message: str, console: bool = True) -> None:
        """Warning message, echoed to the console by default."""
        self._write("WARN", message)
        if console:
            print(message, flush=True)

    def error(self, message: str, console: bool = True) -> None:
        """Error message, echoed to stderr by default."""
        self._write("ERROR", message)
        if console:
            print(message, file=sys.stderr, flush=True)

    def close(self) -> None:
        """Close the log file."""
        if self._handle:
            self._handle.close()
            self._handle = None


# ════════════════════════════════════════════════════════════════════════════
# Exceptions
# ════════════════════════════════════════════════════════════════════════════

class ApifyError(Exception):
    """Base class for Apify-related failures."""


class KeyExhaustedError(ApifyError):
    """Rate limit / quota exceeded for a key - rotate to another key."""


class KeyAuthError(ApifyError):
    """Authentication failure for a key - mark it dead for this session."""


class ActorInputError(ApifyError):
    """Actor rejected the input or does not exist - retrying won't help."""


class ActorRunTimeoutError(ApifyError):
    """Actor run did not finish within the polling timeout."""


class ActorRunFailedError(ApifyError):
    """Actor run finished with FAILED/ABORTED/TIMED-OUT and produced no items."""


class ApiRequestError(ApifyError):
    """Network error that persisted through retries, or an unexpected HTTP status."""


class NoKeysAvailableError(ApifyError):
    """No usable key remains (optionally: no usable key with a LinkedIn cookie)."""

    def __init__(self, message: str, require_cookie: bool = False) -> None:
        """Store whether the failed lookup required a cookie-enabled key."""
        super().__init__(message)
        self.require_cookie = require_cookie


# ════════════════════════════════════════════════════════════════════════════
# Key management
# ════════════════════════════════════════════════════════════════════════════

class ApiKey:
    """Runtime state for a single Apify API token."""

    def __init__(self, token: str, label: str, linkedin_cookie: str | None) -> None:
        """Initialise a key as ACTIVE with zero recorded calls."""
        self.token = token
        self.label = label
        self.linkedin_cookie = linkedin_cookie or None
        self.calls = 0
        self.usage_usd: float | None = None
        self.account_limit_usd: float | None = None
        self.status = "ACTIVE"  # ACTIVE | EXHAUSTED | DEAD
        self.reason = ""
        self.exhausted_at: datetime | None = None

    @property
    def has_cookie(self) -> bool:
        """True when a LinkedIn li_at cookie is configured for this key."""
        return bool(self.li_at)

    @property
    def li_at(self) -> str | None:
        """The bare li_at cookie value (accepts 'li_at=...' or a full cookie header)."""
        if not self.linkedin_cookie:
            return None
        cookie = str(self.linkedin_cookie).strip()
        for part in cookie.split(";"):
            part = part.strip()
            if part.startswith("li_at="):
                return part[len("li_at="):] or None
        return cookie if "=" not in cookie else None


class KeyManager:
    """Round-robin rotation, exhaustion tracking, and budget bookkeeping for API keys."""

    def __init__(self, config: dict, logger: RunLogger, saved_state: dict | None = None) -> None:
        """Build keys from config and restore exhausted keys / call counts from progress."""
        self.logger = logger
        self.max_usage_usd = float(config.get("max_usage_per_key_usd", 4.50))
        self.safety_margin_usd = float(config.get("usage_safety_margin_usd", 0.25))
        self.strategy = str(config.get("rotation_strategy", "round-robin")).lower()
        if self.strategy not in ("round-robin", "sequential"):
            logger.warn(f"⚠ Unknown rotation_strategy '{self.strategy}', using round-robin.")
            self.strategy = "round-robin"
        self.keys: list[ApiKey] = []
        seen: set[str] = set()
        for index, entry in enumerate(config.get("apify_keys", []), start=1):
            label = str(entry.get("label") or f"account-{index}")
            if label in seen:
                label = f"{label}-{index}"
            seen.add(label)
            self.keys.append(ApiKey(str(entry.get("token", "")).strip(), label, entry.get("linkedin_cookie")))
        self._pointer = 0
        self._restore(saved_state or {})

    def _restore(self, saved_state: dict) -> None:
        """Re-apply exhaustion (if < 24h old) and call counts from a previous run."""
        exhausted = saved_state.get("exhausted_keys", {}) or {}
        reasons = saved_state.get("exhausted_reasons", {}) or {}
        calls = saved_state.get("key_calls", {}) or {}
        now = datetime.now()
        for key in self.keys:
            key.calls = int(calls.get(key.label, 0))
            stamp = parse_iso(exhausted.get(key.label))
            if not stamp:
                continue
            if now - stamp < EXHAUSTION_RETRY_AFTER:
                key.status = "EXHAUSTED"
                key.exhausted_at = stamp
                key.reason = reasons.get(key.label, "exhausted in previous run")
                self.logger.info(f"Key {key.label} still exhausted since {stamp.isoformat()} ({key.reason}).")
            else:
                self.logger.info(f"Key {key.label} was exhausted at {stamp.isoformat()} (>24h ago) - retrying it.", console=True)

    def next_key(self, require_cookie: bool = False) -> ApiKey | None:
        """Return the next ACTIVE key per the rotation strategy (None if none qualify)."""
        count = len(self.keys)
        for offset in range(count):
            index = (self._pointer + offset) % count
            key = self.keys[index]
            if key.status != "ACTIVE" or (require_cookie and not key.has_cookie):
                continue
            self._pointer = (index + 1) % count if self.strategy == "round-robin" else index
            return key
        return None

    def mark_exhausted(self, key: ApiKey, reason: str) -> None:
        """Mark a key exhausted (rate limit / quota / budget) with a timestamp."""
        key.status = "EXHAUSTED"
        key.reason = reason
        key.exhausted_at = datetime.now()
        self.logger.warn(f"Key {key.label} marked EXHAUSTED: {reason}", console=False)

    def mark_dead(self, key: ApiKey, reason: str) -> None:
        """Mark a key permanently unusable for this session (authentication failure)."""
        key.status = "DEAD"
        key.reason = reason
        self.logger.error(f"Key {key.label} marked DEAD: {reason}", console=False)

    def record_call(self, key: ApiKey) -> None:
        """Count one actor run started with this key."""
        key.calls += 1

    def update_usage(self, key: ApiKey, usage_usd: float | None, account_limit_usd: float | None) -> None:
        """Store the latest monthly usage numbers reported by Apify."""
        key.usage_usd = usage_usd
        key.account_limit_usd = account_limit_usd

    def budget_for(self, key: ApiKey) -> float:
        """Effective budget: the lower of the configured cap and the account's own limit."""
        if key.account_limit_usd and key.account_limit_usd > 0:
            return min(self.max_usage_usd, key.account_limit_usd)
        return self.max_usage_usd

    def is_over_budget(self, key: ApiKey) -> bool:
        """True when usage is at/near the budget (within the safety margin)."""
        if key.usage_usd is None:
            return False
        return key.usage_usd >= self.budget_for(key) - self.safety_margin_usd

    def available_count(self, require_cookie: bool = False) -> int:
        """Number of ACTIVE keys (optionally only those with a LinkedIn cookie)."""
        return sum(1 for k in self.keys if k.status == "ACTIVE" and (k.has_cookie or not require_cookie))

    def all_exhausted(self) -> bool:
        """True when no ACTIVE key remains."""
        return self.available_count() == 0

    def snapshot(self) -> dict:
        """Serializable state for progress.json (exhausted keys, reasons, call counts)."""
        exhausted = {k.label: k.exhausted_at.isoformat(timespec="seconds")
                     for k in self.keys if k.status == "EXHAUSTED" and k.exhausted_at}
        reasons = {k.label: k.reason for k in self.keys if k.status == "EXHAUSTED"}
        calls = {k.label: k.calls for k in self.keys}
        return {"exhausted_keys": exhausted, "exhausted_reasons": reasons, "key_calls": calls}

    def summary_lines(self) -> list[str]:
        """Human-readable per-key usage lines for the final summary."""
        width = max((len(k.label) for k in self.keys), default=0)
        lines = []
        for key in self.keys:
            used = f"${key.usage_usd:.2f} used" if key.usage_usd is not None else "usage n/a"
            status = key.status + (f" ({key.reason})" if key.status != "ACTIVE" and key.reason else "")
            cookie = " | cookie" if key.has_cookie else ""
            lines.append(f"  {key.label.ljust(width)}: {key.calls} calls | {used} | {status}{cookie}")
        return lines


# ════════════════════════════════════════════════════════════════════════════
# Apify REST client
# ════════════════════════════════════════════════════════════════════════════

class ApifyClient:
    """Thin Apify REST API wrapper with retries, key rotation and run polling."""

    QUOTA_HINTS = ("usage", "quota", "limit-exceeded", "not-enough", "credit",
                   "platform-feature-disabled", "insufficient-funds", "hard-limit")

    def __init__(self, key_manager: KeyManager, logger: RunLogger, actor_timeout: int = 120,
                 poll_interval: int = 5, max_retries: int = 3, http_timeout: int = 30) -> None:
        """Configure timeouts/retries and create a pooled HTTP session."""
        self.keys = key_manager
        self.logger = logger
        self.actor_timeout = actor_timeout
        self.poll_interval = poll_interval
        self.max_retries = max_retries
        self.http_timeout = http_timeout
        self.session = requests.Session()

    # ── low-level HTTP ──────────────────────────────────────────────────────

    @staticmethod
    def _actor_path(actor_id: str) -> str:
        """Apify expects 'username~actor-name' in URL paths."""
        return actor_id.replace("/", "~")

    @staticmethod
    def _parse_error(response: requests.Response) -> tuple[str, str]:
        """Extract (error.type, error.message) from an Apify error response."""
        try:
            error = response.json().get("error", {}) or {}
            return str(error.get("type", "")), str(error.get("message", ""))
        except (ValueError, AttributeError):
            return "", truncate(response.text or "", 300)

    def _is_quota_error(self, status: int, err_type: str, err_msg: str) -> bool:
        """Heuristically detect 'out of credits / monthly limit' errors."""
        if status == 402:
            return True
        lowered_type = err_type.lower()
        if any(hint in lowered_type for hint in self.QUOTA_HINTS):
            return True
        if status == 403:
            lowered_msg = err_msg.lower()
            return any(word in lowered_msg for word in ("usage", "limit", "credit", "quota"))
        return False

    def _backoff(self, attempt: int, what: str) -> None:
        """Sleep 2s, 4s, 8s ... for retry number `attempt` (0-based)."""
        wait = 2 ** (attempt + 1)
        self.logger.warn(f"   ↻ {what} - retrying in {wait}s ({attempt + 1}/{self.max_retries})", console=True)
        time.sleep(wait)

    def _request(self, method: str, path: str, key: ApiKey, params: dict | None = None,
                 json_body: dict | None = None, rotate_on_429: bool = True) -> object:
        """
        Perform an API request with the given key and return the decoded JSON.

        Network errors and 5xx responses are retried with exponential backoff.
        429 raises KeyExhaustedError immediately (or, for polling calls where
        rotating would orphan a running actor, is retried with backoff first).
        """
        query = dict(params or {})
        query["token"] = key.token
        url = API_BASE + path
        for attempt in range(self.max_retries + 1):
            last_attempt = attempt >= self.max_retries
            try:
                response = self.session.request(method, url, params=query, json=json_body, timeout=self.http_timeout)
            except (requests.ConnectionError, requests.Timeout) as exc:
                self.logger.debug(f"[{key.label}] {method} {path} network error: {exc!r}")
                if last_attempt:
                    raise ApiRequestError(f"network error after {self.max_retries} retries: {exc}") from exc
                self._backoff(attempt, f"network error on {method} {path}")
                continue

            status = response.status_code
            self.logger.debug(f"[{key.label}] {method} {path} -> HTTP {status}")
            if 200 <= status < 300:
                if not response.content:
                    return {}
                try:
                    return response.json()
                except ValueError as exc:
                    raise ApiRequestError(f"invalid JSON from {path}") from exc

            err_type, err_msg = self._parse_error(response)
            detail = f"HTTP {status} {err_type} {err_msg}".strip()
            self.logger.debug(f"[{key.label}] error body: {detail}")

            if status == 429:
                if rotate_on_429 or last_attempt:
                    raise KeyExhaustedError("rate limit")
                self._backoff(attempt, f"rate limited while polling {path}")
                continue
            if self._is_quota_error(status, err_type, err_msg):
                raise KeyExhaustedError(f"quota exceeded: {err_type or status}")
            if status in (401, 403):
                raise KeyAuthError(f"auth error: {detail}")
            if status in (400, 404):
                raise ActorInputError(detail)
            if status >= 500 and not last_attempt:
                self._backoff(attempt, f"server error {status} on {path}")
                continue
            raise ApiRequestError(detail)
        raise ApiRequestError(f"request to {path} failed")  # pragma: no cover - loop always returns/raises

    # ── account usage ───────────────────────────────────────────────────────

    def get_usage(self, key: ApiKey) -> tuple[float | None, float | None]:
        """Return (current monthly usage USD, account max monthly usage USD) for a key."""
        payload = self._request("GET", "/users/me/limits", key, rotate_on_429=False)
        data = payload.get("data", {}) if isinstance(payload, dict) else {}
        usage = to_number((data.get("current") or {}).get("monthlyUsageUsd"))
        limit = to_number((data.get("limits") or {}).get("maxMonthlyUsageUsd"))
        return usage, limit

    def acquire_key(self, require_cookie: bool = False) -> ApiKey:
        """
        Pick the next usable key: ACTIVE, has a cookie if required, and under budget.

        Keys found over budget are marked exhausted; keys failing auth are marked
        dead. Raises NoKeysAvailableError when nothing qualifies.
        """
        while True:
            key = self.keys.next_key(require_cookie)
            if key is None:
                scope = "cookie-enabled " if require_cookie else ""
                raise NoKeysAvailableError(f"no {scope}API keys available", require_cookie)
            try:
                usage, limit = self.get_usage(key)
            except KeyAuthError as exc:
                self.keys.mark_dead(key, "auth error")
                self.logger.warn(f"⚠ Key {key.label} rejected ({exc}). Skipping it for this session.")
                continue
            except KeyExhaustedError as exc:
                self.keys.mark_exhausted(key, str(exc))
                self.logger.warn(f"⚠ Key {key.label} exhausted ({exc}) while checking usage.")
                continue
            except ApifyError as exc:
                self.logger.warn(f"   (could not check usage for {key.label}: {exc} - using it anyway)", console=False)
                return key
            self.keys.update_usage(key, usage, limit)
            self.logger.debug(f"[{key.label}] usage ${usage} / budget ${self.keys.budget_for(key):.2f} (account limit {limit})")
            if self.keys.is_over_budget(key):
                reason = f"budget ${usage:.2f}/${self.keys.budget_for(key):.2f}"
                self.keys.mark_exhausted(key, reason)
                self.logger.warn(f"⚠ Key {key.label} exhausted ({reason}). Skipping it.")
                continue
            return key

    # ── actor runs ──────────────────────────────────────────────────────────

    def run_actor(self, actor_id: str, build_input, require_cookie: bool = False,
                  max_items: int | None = None, announce=None) -> list[dict]:
        """
        Run an actor to completion and return its dataset items, rotating keys as needed.

        build_input: dict, or callable(ApiKey) -> dict (e.g. to inject that key's cookie).
        announce:    optional callable(ApiKey) invoked once a key is chosen.
        Raises NoKeysAvailableError, ActorInputError, ActorRunTimeoutError,
        ActorRunFailedError or ApiRequestError.
        """
        failed: tuple[str, str] | None = None
        while True:
            try:
                key = self.acquire_key(require_cookie)
            except NoKeysAvailableError:
                if failed:
                    scope = "cookie-enabled " if require_cookie else ""
                    self.logger.warn(f"⚠ Key {failed[0]} {failed[1]}. No more usable {scope}keys.")
                raise
            if failed:
                self.logger.warn(f"⚠ Key {failed[0]} {failed[1]}. Switching to {key.label}.")
                failed = None
            if announce:
                announce(key)
            payload = build_input(key) if callable(build_input) else build_input
            try:
                return self._run_once(actor_id, payload, key, max_items)
            except KeyExhaustedError as exc:
                self.keys.mark_exhausted(key, str(exc))
                failed = (key.label, f"exhausted ({exc})")
            except KeyAuthError as exc:
                self.keys.mark_dead(key, "auth error")
                self.logger.debug(f"[{key.label}] {exc}")
                failed = (key.label, "dead (auth error)")

    def _run_once(self, actor_id: str, payload: dict, key: ApiKey, max_items: int | None) -> list[dict]:
        """Start one actor run with a specific key, poll until done, and fetch its items."""
        path = self._actor_path(actor_id)
        params: dict = {"timeout": self.actor_timeout + 10}  # server-side safety net
        if max_items:
            params["maxItems"] = max_items  # caps cost for pay-per-result actors
        self.keys.record_call(key)
        self.logger.debug(f"[{key.label}] START {actor_id} input={json.dumps(mask_secrets(payload))}")

        started = self._request("POST", f"/acts/{path}/runs", key, params=params, json_body=payload)
        run = started.get("data", {}) if isinstance(started, dict) else {}
        run_id = run.get("id")
        if not run_id:
            raise ApiRequestError(f"no run id returned for {actor_id}")
        dataset_id = run.get("defaultDatasetId")
        status = run.get("status", "READY")
        self.logger.debug(f"[{key.label}] run {run_id} started (status {status})")

        deadline = time.time() + self.actor_timeout
        while status not in TERMINAL_RUN_STATUSES:
            if time.time() >= deadline:
                self._abort_run(run_id, key)
                raise ActorRunTimeoutError(f"{actor_id} run {run_id} exceeded {self.actor_timeout}s")
            time.sleep(self.poll_interval)
            polled = self._request("GET", f"/acts/{path}/runs/{run_id}", key, rotate_on_429=False)
            run = polled.get("data", {}) if isinstance(polled, dict) else {}
            status = run.get("status", status)
            dataset_id = run.get("defaultDatasetId") or dataset_id
        self.logger.debug(f"[{key.label}] run {run_id} finished with {status}")

        items: list[dict] = []
        if dataset_id:
            fetched = self._request("GET", f"/datasets/{dataset_id}/items", key,
                                    params={"clean": "true", "format": "json"}, rotate_on_429=False)
            if isinstance(fetched, list):
                items = [item for item in fetched if isinstance(item, dict)]
        self.logger.debug(f"[{key.label}] {actor_id} returned {len(items)} item(s): "
                          f"{truncate(json.dumps(mask_secrets(items), default=str))}")
        if status != "SUCCEEDED":
            if not items:
                raise ActorRunFailedError(f"{actor_id} run {run_id} ended with status {status}")
            self.logger.warn(f"   run ended {status} but returned {len(items)} item(s); using them", console=False)
        return items

    def _abort_run(self, run_id: str, key: ApiKey) -> None:
        """Best-effort abort of a run that exceeded the timeout (stops further charges)."""
        try:
            self._request("POST", f"/actor-runs/{run_id}/abort", key, rotate_on_429=False)
            self.logger.debug(f"[{key.label}] aborted run {run_id}")
        except ApifyError as exc:
            self.logger.debug(f"[{key.label}] could not abort run {run_id}: {exc}")


# ════════════════════════════════════════════════════════════════════════════
# Excel I/O
# ════════════════════════════════════════════════════════════════════════════

class ExcelHandler:
    """Reads the contacts sheet, fills empty cells only, and saves checkpoints."""

    def __init__(self, path: str, logger: RunLogger) -> None:
        """Load the workbook's first sheet and map header names to column indexes."""
        self.path = path
        self.logger = logger
        self.workbook = openpyxl.load_workbook(path)
        self.sheet = self.workbook.worksheets[0]
        self.columns = self._map_columns()

    def _map_columns(self) -> dict[str, int]:
        """Find each expected column by header text; fall back to the documented order."""
        headers = {}
        for cell in self.sheet[1]:
            if not is_empty(cell.value):
                headers[str(cell.value).strip().lower()] = cell.column
        columns: dict[str, int] = {}
        missing = []
        for field, header in FIELD_HEADERS.items():
            if header.lower() in headers:
                columns[field] = headers[header.lower()]
            else:
                missing.append(header)
        if not missing:
            return columns
        if len(missing) == len(FIELD_HEADERS):
            self.logger.warn("⚠ No recognised headers in row 1; assuming the documented column order.")
            return {field: index for index, field in enumerate(COLUMN_ORDER, start=1)}
        raise ValueError(f"Missing expected column header(s): {', '.join(missing)}")

    def get(self, row: int, field: str) -> object:
        """Return the raw value of a field in a row."""
        return self.sheet.cell(row=row, column=self.columns[field]).value

    def text(self, row: int, field: str) -> str:
        """Return a field as a stripped string ('' when empty)."""
        value = self.get(row, field)
        return "" if is_empty(value) else str(value).strip()

    def is_blank(self, row: int, field: str) -> bool:
        """True if the cell is empty."""
        return is_empty(self.get(row, field))

    def set_if_empty(self, row: int, field: str, value: object) -> bool:
        """Write value only when the cell is empty; highlight it. Returns True if written."""
        if is_empty(value) or not self.is_blank(row, field):
            return False
        cell = self.sheet.cell(row=row, column=self.columns[field])
        cell.value = value
        if field != "completed":
            cell.fill = HIGHLIGHT_FILL
        return True

    def data_rows(self) -> range:
        """Row numbers below the header."""
        return range(2, self.sheet.max_row + 1)

    def missing_fields(self, row: int) -> list[str]:
        """Enrichable fields that are empty in this row."""
        return [f for f in ("email", "street", "city", "state", "zip", "phone", "website", "revenue", "employees")
                if self.is_blank(row, f)]

    def is_incomplete(self, row: int) -> bool:
        """A row needs enrichment if it has a company and any target field is empty."""
        if self.is_blank(row, "company"):
            return False
        return any(self.is_blank(row, f) for f in TARGET_FIELDS)

    def incomplete_rows(self) -> list[int]:
        """All row numbers that need enrichment, in sheet order."""
        return [row for row in self.data_rows() if self.is_incomplete(row)]

    def save(self, out_path: str) -> None:
        """Save atomically to out_path (write temp file, then rename)."""
        tmp_path = out_path + ".tmp.xlsx"
        self.workbook.save(tmp_path)
        os.replace(tmp_path, out_path)


# ════════════════════════════════════════════════════════════════════════════
# Result parsing
# ════════════════════════════════════════════════════════════════════════════

class ResultParser:
    """Stateless helpers that turn raw actor dataset items into sheet-ready values."""

    @staticmethod
    def name_matches(candidate: object, person_name: str) -> bool:
        """True when the candidate's name contains the person's first and last name."""
        wanted = [t for t in tokenize(person_name) if t not in NAME_SUFFIXES]
        found = set(tokenize(candidate))
        if not wanted or not found:
            return False
        return wanted[0] in found and wanted[-1] in found

    @staticmethod
    def company_overlap(candidate: object, company: str) -> int:
        """Number of significant company-name tokens shared by candidate and company."""
        return len(company_tokens(candidate) & company_tokens(company))

    @staticmethod
    def pick_person(items: list[dict], name: str, company: str, city: str) -> dict | None:
        """
        Choose the best LinkedIn search result for a contact.

        Requires a first+last name match plus either a company or a city match.
        Returns {"linkedin_url", "email", "email_source", "company_match"} or None.
        """
        best, best_score = None, -1
        for item in items:
            if item.get("error") or not ResultParser.name_matches(item.get("name"), name):
                continue
            employer_text = " ".join(str(item.get(k) or "") for k in ("current_company", "headline", "current_title"))
            company_match = ResultParser.company_overlap(employer_text, company) > 0 or not company_tokens(company)
            location_text = str(item.get("location") or "")
            city_match = bool(city) and set(tokenize(city)) <= set(tokenize(location_text))
            if not (company_match or city_match):
                continue
            score = (2 if company_match else 0) + (1 if city_match else 0) + (1 if item.get("email") else 0)
            if score > best_score:
                best_score = score
                best = {"item": item, "company_match": company_match}
        if not best:
            return None
        item = best["item"]
        email = item.get("email")
        if not email and isinstance(item.get("emails"), list) and item["emails"]:
            email = item["emails"][0]
        url = item.get("url") or item.get("profileUrl") or item.get("linkedinUrl")
        if url and "linkedin.com/in/" not in str(url):
            url = None
        return {
            "linkedin_url": str(url).strip() if url else None,
            "email": email.strip() if is_valid_email(email) else None,
            "email_source": item.get("email_source"),
            "company_match": best["company_match"],
        }

    @staticmethod
    def parse_company(items: list[dict], company: str) -> dict | None:
        """
        Extract company fields from Owler-style items (both supported actors).

        Returns a dict with website/street/city/state/zip/revenue/employees/phone
        (values may be None), or None if no item plausibly matches the company.
        """
        wanted = company_tokens(company)
        best, best_overlap = None, -1
        for item in items:
            if item.get("error"):
                continue
            overlap = ResultParser.company_overlap(item.get("name") or item.get("legalName"), company)
            if wanted and item.get("name") and overlap == 0:
                continue
            if overlap > best_overlap:
                best, best_overlap = item, overlap
        if best is None:
            return None

        address = best.get("address") if isinstance(best.get("address"), dict) else {}
        street = " ".join(str(address.get(k) or "").strip() for k in ("street1",)).strip()
        if address.get("street2"):
            street = f"{street}, {str(address['street2']).strip()}" if street else str(address["street2"]).strip()
        city = address.get("city")
        state = address.get("state")
        if (not city or not state) and isinstance(best.get("headquarters"), str):
            parts = [p.strip() for p in best["headquarters"].split(",")]
            city = city or (parts[0] if parts else None)
            state = state or (parts[1] if len(parts) > 1 else None)

        revenue = None
        amount = to_number(best.get("revenueAmount")) or to_number(best.get("revenueEstimateUsd"))
        if amount:
            revenue = format_revenue(amount)
        elif not is_empty(best.get("revenueFormatted")):
            text = str(best["revenueFormatted"]).strip()
            revenue = text if text.startswith("$") else f"${text}"

        employees = to_number(best.get("employees"))
        if employees is None:
            employees = to_number(best.get("employeeEstimate"))

        website = best.get("website")
        if is_empty(website) and not is_empty(best.get("domain")):
            website = f"https://{str(best['domain']).strip()}/"

        return {
            "name": best.get("name"),
            "website": str(website).strip() if not is_empty(website) else None,
            "street": street or None,
            "city": str(city).strip() if not is_empty(city) else None,
            "state": state_abbrev(state),
            "zip": str(address.get("zipcode")).strip() if not is_empty(address.get("zipcode")) else None,
            "revenue": revenue,
            "employees": int(employees) if employees else None,
            "phone": normalize_phone(best.get("phoneNumber") or best.get("phone")),
        }

    @staticmethod
    def find_phone(data: object) -> str | None:
        """Recursively search actor output for the first plausible phone number."""
        if isinstance(data, dict):
            for key, value in data.items():
                lowered = str(key).lower()
                if "phone" in lowered or "mobile" in lowered:
                    candidates = value if isinstance(value, list) else [value]
                    for candidate in candidates:
                        if isinstance(candidate, dict):
                            candidate = candidate.get("number") or candidate.get("value") or candidate.get("phone")
                        if isinstance(candidate, (str, int)) and not isinstance(candidate, bool):
                            phone = normalize_phone(candidate)
                            if phone:
                                return phone
            for value in data.values():
                if isinstance(value, (dict, list)):
                    phone = ResultParser.find_phone(value)
                    if phone:
                        return phone
        elif isinstance(data, list):
            for value in data:
                phone = ResultParser.find_phone(value)
                if phone:
                    return phone
        return None


# ════════════════════════════════════════════════════════════════════════════
# Orchestration
# ════════════════════════════════════════════════════════════════════════════

STAT_KEYS = [
    "emails_found", "phones_found", "websites_found", "addresses_found", "revenue_found",
    "employees_found", "linkedin_urls_found", "rows_enriched", "contacts_failed", "total_processed",
]


class Enricher:
    """Main orchestration: selects rows, runs the actor pipeline, merges results, checkpoints."""

    def __init__(self, args: argparse.Namespace, logger: RunLogger, paths: dict[str, str]) -> None:
        """Store CLI options and file paths; heavy setup happens in run()/dry_run()/show_status()."""
        self.args = args
        self.logger = logger
        self.paths = paths
        self.config: dict = {}
        self.actors: dict[str, str] = dict(DEFAULT_ACTORS)
        self.excel: ExcelHandler | None = None
        self.key_manager: KeyManager | None = None
        self.client: ApifyClient | None = None
        self.progress: dict = {}
        self.disabled_actors: set[str] = set()
        self.phones_enabled = not args.skip_phones
        self.phone_skip_logged = False
        self._lookups_announced = 0  # lookups started for the current contact (console bookkeeping)
        self.delay = 5.0
        self.batch_size = 10

    # ── configuration & progress ────────────────────────────────────────────

    def load_config(self) -> dict:
        """Load and validate keys.json; raises ValueError/OSError with a clear message."""
        path = self.paths["keys"]
        if not os.path.isfile(path):
            raise ValueError(f"Keys file not found: {path} (copy the keys.json template and add your tokens)")
        try:
            config = load_json_file(path)
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSON in {path}: {exc}") from exc
        keys = config.get("apify_keys")
        if not isinstance(keys, list) or not keys:
            raise ValueError(f"{path}: 'apify_keys' must be a non-empty list")
        for index, entry in enumerate(keys, start=1):
            token = str(entry.get("token", "")).strip()
            label = entry.get("label") or f"#{index}"
            if not token or "xxx" in token.lower() or "replace" in token.lower():
                raise ValueError(f"{path}: key {label} has a missing or placeholder token")
            cookie = entry.get("linkedin_cookie")
            if cookie and ("replace" in str(cookie).lower() or "xxx" in str(cookie).lower()):
                self.logger.warn(f"⚠ Key {label}: linkedin_cookie looks like a placeholder - ignoring it.")
                entry["linkedin_cookie"] = None
        if int(config.get("batch_size", 10)) < 1:
            raise ValueError(f"{path}: batch_size must be >= 1")
        if float(config.get("delay_between_calls_seconds", 5)) < 0:
            raise ValueError(f"{path}: delay_between_calls_seconds must be >= 0")
        self.actors.update({k: v for k, v in (config.get("actors") or {}).items() if v})
        return config

    def load_progress(self) -> dict | None:
        """Read progress.json if present; returns None when missing or unreadable."""
        path = self.paths["progress"]
        if not os.path.isfile(path):
            return None
        try:
            with open(path, "r", encoding="utf-8") as handle:
                data = json.load(handle)
        except (OSError, json.JSONDecodeError) as exc:
            self.logger.warn(f"⚠ Could not read {path}: {exc}")
            return None
        data.setdefault("stats", {})
        for stat in STAT_KEYS:
            data["stats"].setdefault(stat, 0)
        data.setdefault("company_cache", {})
        data.setdefault("failed_rows", [])
        return data

    def new_progress(self, total: int) -> dict:
        """Fresh progress structure for a new run."""
        return {
            "input_file": os.path.basename(self.paths["input"]),
            "last_processed_row": 0,
            "exhausted_keys": {},
            "exhausted_reasons": {},
            "key_calls": {},
            "started_at": now_iso(),
            "updated_at": now_iso(),
            "total_incomplete": total,
            "stats": {stat: 0 for stat in STAT_KEYS},
            "company_cache": {},
            "failed_rows": [],
        }

    @property
    def stats(self) -> dict:
        """Shortcut to the running statistics dict inside progress."""
        return self.progress["stats"]

    # ── main run ────────────────────────────────────────────────────────────

    def run(self) -> int:
        """Execute enrichment. Returns a process exit code (0 ok, 1 error, 2 keys exhausted, 130 interrupted)."""
        try:
            self.config = self.load_config()
        except (OSError, ValueError) as exc:
            self.logger.error(f"✖ {exc}")
            return 1

        saved = self.load_progress()
        resume = self.args.resume
        if not resume and saved and int(saved.get("last_processed_row", 0)) > 0 and not self.args.restart:
            self.logger.error(
                f"✖ {self.paths['progress']} shows an earlier run (last row {saved['last_processed_row']}).\n"
                "  Use --resume to continue it, or --restart to start over (overwrites the enriched file).")
            return 1
        if resume and not saved:
            self.logger.warn("⚠ --resume given but no progress.json found - starting a fresh run.")
            resume = False
        if resume and saved.get("input_file") not in (None, os.path.basename(self.paths["input"])):
            self.logger.warn(f"⚠ progress.json was created for '{saved.get('input_file')}', not this file.")

        source = self.paths["output"] if resume and os.path.isfile(self.paths["output"]) else self.paths["input"]
        try:
            self.excel = ExcelHandler(source, self.logger)
        except (OSError, ValueError) as exc:
            self.logger.error(f"✖ Could not open {source}: {exc}")
            return 1

        if resume:
            self.progress = saved
            start_after = int(saved.get("last_processed_row", 0))
            key_state = saved
        else:
            self.progress = self.new_progress(len(self.excel.incomplete_rows()))
            start_after = 0
            key_state = {k: saved.get(k, {}) for k in ("exhausted_keys", "exhausted_reasons")} if saved else None

        self.key_manager = KeyManager(self.config, self.logger, key_state)
        self.client = ApifyClient(
            self.key_manager, self.logger,
            actor_timeout=int(self.config.get("actor_timeout_seconds", 120)),
            poll_interval=int(self.config.get("poll_interval_seconds", 5)),
        )
        self.delay = float(self.args.delay if self.args.delay is not None
                           else self.config.get("delay_between_calls_seconds", 5))
        self.batch_size = int(self.config.get("batch_size", 10))

        if self.phones_enabled and self.key_manager.available_count(require_cookie=True) == 0:
            self.logger.warn("⚠ No keys with a linkedin_cookie are available - phone scraping will be skipped.")
            self.phone_skip_logged = True

        pending = [row for row in self.excel.incomplete_rows() if row > start_after]
        if self.args.test:
            pending = pending[: self.args.test]
        total = int(self.progress.get("total_incomplete") or len(pending))

        self.logger.info(
            f"Source: {os.path.basename(source)} | to process now: {len(pending)} | "
            f"already processed: {self.stats['total_processed']}/{total} | keys: {len(self.key_manager.keys)} "
            f"({self.key_manager.available_count()} active, {self.key_manager.available_count(True)} with cookie) | "
            f"delay {self.delay:g}s | checkpoint every {self.batch_size}", console=True)
        self.logger.info(f"Log: {self.paths['log']}\n", console=True)

        if self.key_manager.all_exhausted():
            return self._stop_all_exhausted(total)
        if not pending:
            self.logger.info("Nothing left to enrich.", console=True)
            self.checkpoint()
            self.print_summary(total)
            return 0

        since_checkpoint = 0
        try:
            for row in pending:
                number = self.stats["total_processed"] + 1
                try:
                    self.process_contact(row, number, total)
                except NoKeysAvailableError:
                    return self._stop_all_exhausted(total)
                except Exception as exc:  # never let one contact crash the run
                    self.stats["contacts_failed"] += 1
                    self.progress["failed_rows"].append({"row": row, "error": repr(exc), "at": now_iso()})
                    self.logger.error(f"   ✖ Row {row} failed: {exc!r} - continuing")
                self.progress["last_processed_row"] = row
                self.stats["total_processed"] += 1
                since_checkpoint += 1
                if since_checkpoint >= self.batch_size:
                    self.checkpoint()
                    since_checkpoint = 0
        except KeyboardInterrupt:
            self.logger.warn("\n⚠ Interrupted - saving progress (the current contact will be retried on --resume)...")
            self.checkpoint()
            self.print_summary(total)
            return 130

        self.checkpoint()
        self.print_summary(total)
        return 0

    def process_contact(self, row: int, number: int, total: int) -> None:
        """Run the 3-step actor pipeline for one row and merge results into empty cells."""
        excel = self.excel
        name, company = excel.text(row, "name"), excel.text(row, "company")
        city = excel.text(row, "city")
        missing = excel.missing_fields(row)
        prefix = f"Processing {number}/{total}..."
        label = f"{name or '(no name)'} @ {company}"
        self.logger.info(f"── Row {row}: {label} | missing: {', '.join(missing)}")
        filled: list[str] = []
        self._lookups_announced = 0
        phone_wanted = self.phones_enabled and "phone" in missing

        # Step 1 - LinkedIn search (email + profile URL)
        linkedin_url = None
        wants_url = phone_wanted and self.key_manager.available_count(require_cookie=True) > 0
        if not name:
            self.logger.info(f"{prefix} Row {row} has no name - skipping LinkedIn search", console=True)
        elif "linkedin_search" not in self.disabled_actors and ("email" in missing or wants_url):
            person = self._search_person(name, company, city, prefix)
            if person:
                linkedin_url = person["linkedin_url"]
                if linkedin_url:
                    self.stats["linkedin_urls_found"] += 1
                    self.logger.info(f"   LinkedIn: {linkedin_url}")
                email = person["email"]
                accept_guessed = bool(self.config.get("accept_guessed_emails", True))
                if email and person["company_match"] and (accept_guessed or person["email_source"] != "guessed"):
                    self._fill(row, "email", email, filled)
                elif email:
                    self.logger.info(f"   Email {email} ignored (source={person['email_source']}, "
                                     f"company match={person['company_match']})")

        # Step 2 - Company data (cached per company)
        company_data = None
        if any(field in missing for field in COMPANY_FIELDS):
            company_data = self._lookup_company(company, prefix)
            if company_data:
                self._merge_company(row, company_data, filled)

        # Step 3 - Phone (needs LinkedIn URL + cookie-enabled key)
        if phone_wanted and excel.is_blank(row, "phone"):
            if linkedin_url and "phone" not in self.disabled_actors:
                phone = self._find_phone(linkedin_url, label, prefix)
                if phone:
                    self._fill(row, "phone", phone, filled)
            if excel.is_blank(row, "phone") and company_data and company_data.get("phone"):
                if self._fill(row, "phone", company_data["phone"], filled):
                    self.logger.info("   (phone taken from company main line)")

        if filled:
            self.stats["rows_enriched"] += 1
            excel.set_if_empty(row, "completed", AUTO_ENRICHED_LABEL)
            values = ", ".join(f"{f}={excel.get(row, f)}" for f in filled)
            self.logger.info(f"   ✓ Filled {len(filled)} field(s): {values}", console=True)
        elif self._lookups_announced == 0:
            self.logger.info(f"{prefix} {label} – nothing to look up (missing: {', '.join(missing)})", console=True)
        else:
            self.logger.info("   – No new data found", console=True)

    # ── pipeline steps ──────────────────────────────────────────────────────

    def _run_step(self, step: str, actor_id: str, build_input, prefix: str, action: str,
                  max_items: int | None = None, require_cookie: bool = False,
                  quiet_failure: bool = False) -> list[dict] | None:
        """
        Run one actor with rotation and step-level error handling.

        Returns items (possibly empty) on success, None if the step failed.
        quiet_failure logs FAILED runs to the file only (used when a fallback follows).
        NoKeysAvailableError is propagated so the caller can decide what to do.
        """
        attempted: list[bool] = []

        def announce(key: ApiKey) -> None:
            attempted.append(True)
            self._lookups_announced += 1
            self.logger.info(f"{prefix} [{key.label}] {action}", console=True)

        try:
            return self.client.run_actor(actor_id, build_input, require_cookie=require_cookie,
                                         max_items=max_items, announce=announce)
        except ActorInputError as exc:
            self.disabled_actors.add(step)
            self.logger.warn(f"⚠ {actor_id} rejected the request ({exc}). Disabling step '{step}' for this run "
                             "- check the actor ID / input in keys.json.")
        except ActorRunTimeoutError as exc:
            self.logger.warn(f"   ⏱ Timeout: {exc} - skipping this step")
        except ActorRunFailedError as exc:
            self.logger.warn(f"   ✖ {exc} - skipping this step", console=not quiet_failure)
        except ApiRequestError as exc:
            self.logger.warn(f"   ✖ {exc} - skipping this step")
        finally:
            if attempted and self.delay > 0:
                time.sleep(self.delay)
        return None

    def _search_person(self, name: str, company: str, city: str, prefix: str) -> dict | None:
        """Actor 1: search LinkedIn for 'Name Company' and pick the matching profile."""
        payload = {
            "mode": "search_profiles",
            "searchQuery": f"{name} {company}",
            "maxResults": 5,
            "discoverEmails": True,
        }
        items = self._run_step("linkedin_search", self.actors["linkedin_search"], payload, prefix,
                               f"Searching: {name} @ {company}", max_items=5)
        if items is None:
            return None
        if not items:
            self.logger.info("   LinkedIn search: no results")
            return None
        person = ResultParser.pick_person(items, name, company, city)
        if not person:
            self.logger.info(f"   LinkedIn search: {len(items)} result(s), none matched {name} / {company}")
        return person

    def _lookup_company(self, company: str, prefix: str) -> dict | None:
        """Actor 2: Owler company data (primary, then fallback), cached per company name."""
        cache = self.progress["company_cache"]
        cache_key = company_cache_key(company)
        if cache_key in cache:
            self._lookups_announced += 1
            self.logger.info(f"{prefix} [cache] Company: {company}", console=True)
            return cache[cache_key] or None

        data, had_error = None, False
        common_flags = {"includeCompetitors": False, "includeFunding": False, "includeAcquisitions": False,
                        "includeInvestments": False, "includeLeadership": False}
        slug = owler_slug(company)
        if slug and "company_primary" not in self.disabled_actors:
            payload = {"companyUrls": [{"url": f"https://www.owler.com/company/{slug}"}], "maxResults": 1,
                       "includeNewsEvents": False, "includeFinancialHistory": False, "includeSocial": False,
                       **common_flags}
            items = self._run_step("company_primary", self.actors["company_primary"], payload, prefix,
                                   f"Company: {company}", max_items=1, quiet_failure=True)
            if items is None:
                had_error = True
            else:
                data = ResultParser.parse_company(items, company)

        if not data and "company_fallback" not in self.disabled_actors:
            payload = {"companyNames": [company], "maxItems": 1, "maxConcurrency": 1, "includeNews": False,
                       **common_flags}
            items = self._run_step("company_fallback", self.actors["company_fallback"], payload, prefix,
                                   f"Company (fallback): {company}", max_items=1)
            if items is None:
                had_error = True
            else:
                data = ResultParser.parse_company(items, company)
                had_error = False  # fallback answered definitively

        if data:
            self.logger.info(f"   Company match: {data.get('name')} | {json.dumps(data, default=str)}")
        else:
            self.logger.info(f"   Company data not found for {company}")
        if data or not had_error:
            cache[cache_key] = data or {}
        return data

    def _find_phone(self, linkedin_url: str, label: str, prefix: str) -> str | None:
        """Actor 3: phone number from a LinkedIn profile, using a cookie-enabled key."""
        if self.key_manager.available_count(require_cookie=True) == 0:
            if not self.phone_skip_logged:
                self.logger.warn("⚠ No cookie-enabled keys available - skipping phone enrichment.")
                self.phone_skip_logged = True
            return None
        template = self.config.get("phone_actor_input") or DEFAULT_PHONE_INPUT_TEMPLATE

        def build_input(key: ApiKey) -> dict:
            return self._fill_template(template, {"{profile_url}": linkedin_url, "{li_at}": key.li_at or ""})

        try:
            items = self._run_step("phone", self.actors["phone"], build_input, prefix,
                                   f"Phone: {label}", require_cookie=True)
        except NoKeysAvailableError:
            if self.key_manager.all_exhausted():
                raise
            if not self.phone_skip_logged:
                self.logger.warn("⚠ Cookie-enabled keys exhausted - skipping phone enrichment from now on.")
                self.phone_skip_logged = True
            return None
        if not items:
            if items is not None:
                self.logger.info("   Phone scraper: no results")
            return None
        phone = ResultParser.find_phone(items)
        if not phone:
            self.logger.info("   Phone scraper: results contained no phone number")
        return phone

    @staticmethod
    def _fill_template(template: object, values: dict[str, str]) -> object:
        """Recursively substitute placeholders like '{profile_url}' in a JSON template."""
        if isinstance(template, dict):
            return {k: Enricher._fill_template(v, values) for k, v in template.items()}
        if isinstance(template, list):
            return [Enricher._fill_template(v, values) for v in template]
        if isinstance(template, str):
            for placeholder, value in values.items():
                template = template.replace(placeholder, value)
        return template

    # ── merging ─────────────────────────────────────────────────────────────

    def _fill(self, row: int, field: str, value: object, filled: list[str]) -> bool:
        """Fill one empty cell, update stats, and record the field name. Returns True if written."""
        if not self.excel.set_if_empty(row, field, value):
            return False
        filled.append(field)
        stat = FIELD_STAT.get(field)
        if stat:
            self.stats[stat] += 1
        return True

    def _merge_company(self, row: int, data: dict, filled: list[str]) -> None:
        """Merge company data; address is only used when the HQ city matches the contact's city."""
        for field in ("website", "revenue", "employees"):
            self._fill(row, field, data.get(field), filled)
        row_city = self.excel.text(row, "city")
        if row_city and data.get("city") and tokenize(row_city) != tokenize(data["city"]):
            self.logger.info(f"   HQ city '{data['city']}' differs from contact city '{row_city}' - address not used")
            return
        address_found = False
        for field in ("street", "zip"):
            if self._fill(row, field, data.get(field), filled):
                address_found = True
        for field in ("city", "state"):
            self._fill(row, field, data.get(field), filled)
        if address_found:
            self.stats["addresses_found"] += 1

    # ── checkpoints & reporting ─────────────────────────────────────────────

    def checkpoint(self) -> None:
        """Save the enriched workbook and progress.json (errors are logged, not raised)."""
        self.progress.update(self.key_manager.snapshot())
        self.progress["updated_at"] = now_iso()
        try:
            self.excel.save(self.paths["output"])
            atomic_write_json(self.paths["progress"], self.progress)
            self.logger.info(f"💾 Checkpoint saved ({os.path.basename(self.paths['output'])}, "
                             f"last row {self.progress['last_processed_row']})", console=True)
        except OSError as exc:
            self.logger.error(f"✖ Checkpoint failed ({exc}). Is the output file open in Excel? Will retry.")

    def _stop_all_exhausted(self, total: int) -> int:
        """Save everything and exit gracefully when no key can be used."""
        self.checkpoint()
        self.print_summary(total)
        message = (f"All API keys exhausted. {self.stats['rows_enriched']}/{total} contacts enriched. "
                   "Run again with fresh keys using --resume")
        self.logger.warn(f"\n{message}")
        return 2

    def print_summary(self, total: int) -> None:
        """Print (and log) the end-of-run summary table."""
        s = self.stats
        lines = [
            "",
            "═══ ENRICHMENT SUMMARY ═══",
            f"Contacts processed: {s['total_processed']} / {total}",
            f"Contacts enriched:  {s['rows_enriched']}",
            f"Emails found:       {s['emails_found']}",
            f"Phones found:       {s['phones_found']}",
            f"Websites found:     {s['websites_found']}",
            f"Addresses found:    {s['addresses_found']}",
            f"Revenue found:      {s['revenue_found']}",
            f"Employees found:    {s['employees_found']}",
            f"LinkedIn matches:   {s['linkedin_urls_found']}",
            f"Failed contacts:    {s['contacts_failed']}",
        ]
        if self.disabled_actors:
            lines.append(f"Disabled steps:     {', '.join(sorted(self.disabled_actors))}")
        lines += ["", "Key Usage:", *self.key_manager.summary_lines(), "",
                  f"Output: {os.path.basename(self.paths['output'])}",
                  f"Log: {os.path.basename(self.paths['log'])}"]
        text = "\n".join(lines)
        self.logger.info(text, console=True)

    # ── --dry-run ───────────────────────────────────────────────────────────

    def dry_run(self) -> int:
        """List what would be enriched and estimate actor runs, without calling any API."""
        saved = self.load_progress() if self.args.resume else None
        source = self.paths["output"] if saved and os.path.isfile(self.paths["output"]) else self.paths["input"]
        try:
            excel = ExcelHandler(source, self.logger)
        except (OSError, ValueError) as exc:
            self.logger.error(f"✖ Could not open {source}: {exc}")
            return 1
        start_after = int(saved.get("last_processed_row", 0)) if saved else 0
        rows = [row for row in excel.incomplete_rows() if row > start_after]
        if self.args.test:
            rows = rows[: self.args.test]

        cookie_keys, key_note = 0, ""
        try:
            config = self.load_config()
            cookie_keys = sum(1 for k in config["apify_keys"] if ApiKey("", "", k.get("linkedin_cookie")).has_cookie)
            key_note = f"{len(config['apify_keys'])} key(s) configured, {cookie_keys} with a LinkedIn cookie"
        except (OSError, ValueError) as exc:
            key_note = f"keys file problem: {exc}"
        phones = self.phones_enabled and cookie_keys > 0

        print(f"DRY RUN - {os.path.basename(source)}: {len(rows)} contact(s) would be processed\n")
        companies: set[str] = set()
        searches = company_lookups = phone_lookups = 0
        for index, row in enumerate(rows, start=1):
            name, company = excel.text(row, "name"), excel.text(row, "company")
            missing = excel.missing_fields(row)
            plan = []
            phone_wanted = phones and "phone" in missing
            if name and ("email" in missing or phone_wanted):
                plan.append("linkedin-search")
                searches += 1
            if any(f in missing for f in COMPANY_FIELDS):
                key = company_cache_key(company)
                if key in companies:
                    plan.append("company (cached)")
                else:
                    companies.add(key)
                    plan.append("company")
                    company_lookups += 1
            if phone_wanted and name:
                plan.append("phone (if LinkedIn URL found)")
                phone_lookups += 1
            location = ", ".join(p for p in (excel.text(row, "city"), excel.text(row, "state")) if p)
            print(f"{index:>4}. Row {row:<4} {name or '(no name)'} @ {company} ({location})")
            print(f"       missing: {', '.join(missing)}")
            print(f"       plan:    {', '.join(plan) or 'nothing (no name)'}")

        print("\nEstimated actor runs (maximum):")
        print(f"  LinkedIn searches: {searches}")
        print(f"  Company lookups:   {company_lookups} unique companies (up to 2 runs each with fallback)")
        print(f"  Phone lookups:     {phone_lookups if phones else 0}"
              + ("" if phones else "  (disabled: --skip-phones or no cookie keys)"))
        print(f"\nKeys: {key_note}")
        print(f"Output would be: {os.path.basename(self.paths['output'])}")
        print("No API calls were made.")
        return 0

    # ── --status ────────────────────────────────────────────────────────────

    def show_status(self) -> int:
        """Show saved progress and live usage for every configured key (read-only API calls)."""
        saved = self.load_progress()
        print("═══ STATUS ═══")
        if saved:
            s = saved["stats"]
            print(f"Input file:        {saved.get('input_file')}")
            print(f"Started:           {saved.get('started_at')}   Updated: {saved.get('updated_at')}")
            print(f"Last row:          {saved.get('last_processed_row')}")
            print(f"Processed:         {s['total_processed']} / {saved.get('total_incomplete')}"
                  f"   (enriched {s['rows_enriched']}, failed {s['contacts_failed']})")
            print(f"Found:             emails {s['emails_found']}, phones {s['phones_found']}, "
                  f"websites {s['websites_found']}, addresses {s['addresses_found']}, "
                  f"revenue {s['revenue_found']}, employees {s['employees_found']}")
            print(f"Companies cached:  {len(saved.get('company_cache', {}))}")
        else:
            print("No progress.json yet - nothing has been processed.")

        source = self.paths["output"] if os.path.isfile(self.paths["output"]) else self.paths["input"]
        try:
            remaining = len(ExcelHandler(source, self.logger).incomplete_rows())
            print(f"Incomplete rows in {os.path.basename(source)}: {remaining}")
        except (OSError, ValueError) as exc:
            print(f"Could not read {source}: {exc}")

        try:
            config = self.load_config()
        except (OSError, ValueError) as exc:
            print(f"\nKeys: {exc}")
            return 1
        manager = KeyManager(config, self.logger, saved)
        client = ApifyClient(manager, self.logger, max_retries=1)
        print("\nKey Usage (live):")
        for key in manager.keys:
            cookie = "cookie" if key.has_cookie else "no cookie"
            state = key.status + (f" since {key.exhausted_at.isoformat(timespec='minutes')} ({key.reason})"
                                  if key.exhausted_at else "")
            try:
                usage, limit = client.get_usage(key)
                manager.update_usage(key, usage, limit)
                usage_text = (f"${usage or 0:.2f} / ${manager.budget_for(key):.2f} budget"
                              + (f" (account limit ${limit:.2f})" if limit else ""))
                if manager.is_over_budget(key):
                    usage_text += " OVER BUDGET"
            except KeyAuthError:
                usage_text = "INVALID TOKEN"
            except ApifyError as exc:
                usage_text = f"usage unavailable ({exc})"
            print(f"  {key.label}: {key.calls} calls | {usage_text} | {state} | {cookie}")
        return 0


# ════════════════════════════════════════════════════════════════════════════
# CLI
# ════════════════════════════════════════════════════════════════════════════

def build_arg_parser() -> argparse.ArgumentParser:
    """Define the command-line interface."""
    parser = argparse.ArgumentParser(
        description="Enrich business contacts in an Excel file using Apify actors (multi-key, resumable).")
    parser.add_argument("input", help="Input .xlsx file")
    parser.add_argument("--dry-run", action="store_true", help="Show what would be enriched; no API calls")
    parser.add_argument("--test", type=int, metavar="N", help="Process only the first N incomplete contacts")
    parser.add_argument("--resume", action="store_true", help="Resume from progress.json")
    parser.add_argument("--restart", action="store_true",
                        help="Ignore existing progress.json and start over (overwrites the enriched file)")
    parser.add_argument("--status", action="store_true", help="Show progress and live key usage")
    parser.add_argument("--skip-phones", action="store_true", help="Skip phone enrichment entirely")
    parser.add_argument("--keys", default="keys.json", help="Keys/config file (default: keys.json)")
    parser.add_argument("--delay", type=float, metavar="SECONDS", help="Override delay between actor calls")
    return parser


def main(argv: list[str] | None = None) -> int:
    """Parse arguments, resolve paths, and dispatch to run / dry-run / status."""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")  # avoid crashes on consoles without UTF-8
        except (AttributeError, ValueError):
            pass

    parser = build_arg_parser()
    args = parser.parse_args(argv)
    if args.test is not None and args.test < 1:
        parser.error("--test must be >= 1")
    if args.delay is not None and args.delay < 0:
        parser.error("--delay must be >= 0")
    if args.resume and args.restart:
        parser.error("--resume and --restart are mutually exclusive")

    input_path = os.path.abspath(args.input)
    if not os.path.isfile(input_path):
        print(f"✖ Input file not found: {args.input}", file=sys.stderr)
        return 1
    base_dir = os.path.dirname(input_path)
    stem = os.path.splitext(os.path.basename(input_path))[0].strip().replace(" ", "_")
    keys_path = args.keys
    if not os.path.isfile(keys_path) and not os.path.isabs(keys_path):
        beside_input = os.path.join(base_dir, keys_path)
        keys_path = beside_input if os.path.isfile(beside_input) else keys_path

    paths = {
        "input": input_path,
        "output": os.path.join(base_dir, f"enriched_{stem}.xlsx"),
        "progress": os.path.join(base_dir, "progress.json"),
        "log": os.path.join(base_dir, f"enrichment_log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt"),
        "keys": keys_path,
    }
    is_real_run = not (args.dry_run or args.status)
    logger = RunLogger(paths["log"] if is_real_run else None)
    enricher = Enricher(args, logger, paths)
    try:
        if args.status:
            return enricher.show_status()
        if args.dry_run:
            return enricher.dry_run()
        logger.info(f"Run started with args: {vars(args)}")
        return enricher.run()
    except KeyboardInterrupt:
        print("\nInterrupted.", file=sys.stderr)
        return 130
    finally:
        logger.close()


if __name__ == "__main__":
    sys.exit(main())
